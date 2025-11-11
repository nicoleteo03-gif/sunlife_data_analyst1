import pdfplumber
from docx import Document
import pytesseract
from PIL import Image
from io import BytesIO
import os
import re

def extract_text_from_docx(path):
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)

def extract_text_from_pdf(path):
    texts = []
    tables = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                try:
                    text = page.extract_text() or ""
                    if text.strip():
                        texts.append(text)
                    # table extraction (basic)
                    page_tables = page.extract_tables()
                    for t in page_tables:
                        # convert table rows to string lines
                        rows = [" | ".join([cell or "" for cell in row]) for row in t]
                        tables.append("\n".join(rows))
                except Exception:
                    continue
    except Exception as e:
        print("pdfplumber error:", e)
    return "\n\n".join(texts), tables

def extract_tables_from_pdf(path):
    """
    Return a list of table strings extracted from a PDF (same format as extract_text_from_pdf returns for tables).
    This makes ingest.py imports safe if it imports this function directly.
    """
    tables = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                try:
                    page_tables = page.extract_tables()
                    for t in page_tables:
                        rows = [" | ".join([cell or "" for cell in row]) for row in t]
                        tables.append("\n".join(rows))
                except Exception:
                    continue
    except Exception as e:
        print("pdfplumber error (tables):", e)
    return tables

def ocr_image(path):
    try:
        img = Image.open(path)
        text = pytesseract.image_to_string(img)
        return text
    except Exception as e:
        print("OCR error", e)
        return ""

def extract_document(path):
    """
    Returns (text, tables, metadata)
    - text: extracted text (string)
    - tables: list of table strings (may be empty)
    - metadata: dict with file_ext and size_bytes
    """
    path = str(path)
    ext = os.path.splitext(path)[1].lower()
    text = ""
    tables = []
    metadata = {}
    if ext in [".pdf"]:
        t, tbls = extract_text_from_pdf(path)
        text = t
        tables = tbls
        if not text.strip() and not tables:
            # scanned PDF fallback: try per-page OCR (simple, may be slow)
            try:
                # attempt simple OCR on each page by rendering page images if available
                with pdfplumber.open(path) as pdf:
                    ocr_texts = []
                    for page in pdf.pages:
                        try:
                            pil_image = page.to_image(resolution=150).original
                            ocr_texts.append(pytesseract.image_to_string(pil_image))
                        except Exception:
                            continue
                    text = "\n".join([t for t in ocr_texts if t.strip()])
            except Exception:
                text = ""
    elif ext in [".docx"]:
        text = extract_text_from_docx(path)
    elif ext in [".png", ".jpg", ".jpeg", ".tiff", ".bmp"]:
        text = ocr_image(path)
    else:
        # fallback: try reading as text
        try:
            with open(path, "r", encoding="utf-8") as f:
                text = f.read()
        except Exception:
            text = ""

    metadata["file_ext"] = ext
    try:
        metadata["size_bytes"] = os.path.getsize(path)
    except Exception:
        metadata["size_bytes"] = 0
    return text, tables, metadata